import streamlit as st
import sqlite3
from pathlib import Path
import pandas as pd
from datetime import date
import base64


# =========================================================
# DATABASE
# =========================================================

DB_PATH = "koc_data.db"

conn = sqlite3.connect(
    DB_PATH,
    check_same_thread=False
)

cursor = conn.cursor()


cursor.execute("""
CREATE TABLE IF NOT EXISTS koc (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    username TEXT UNIQUE,
    phone TEXT,
    name TEXT,
    follower INTEGER DEFAULT 0,
    category TEXT,
    note TEXT
)
""")


cursor.execute("""
CREATE TABLE IF NOT EXISTS brands (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    brand_name TEXT UNIQUE,
    commission REAL DEFAULT 0,
    target_gmv REAL DEFAULT 0,
    note TEXT
)
""")


cursor.execute("""
CREATE TABLE IF NOT EXISTS analytics (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    report_date TEXT,
    report_month TEXT,
    creator_name TEXT,
    store_name TEXT,
    gmv REAL DEFAULT 0,
    ac REAL DEFAULT 0
)
""")


cursor.execute("""
CREATE TABLE IF NOT EXISTS tap_targets (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    target_month TEXT UNIQUE,
    target_ac REAL DEFAULT 0,
    note TEXT
)
""")




cursor.execute("""
CREATE TABLE IF NOT EXISTS booking_services (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    contract_month TEXT,
    brand_name TEXT,
    service_group TEXT,
    package_name TEXT,
    tier TEXT,
    contract_fee REAL DEFAULT 0,
    running INTEGER DEFAULT 0,
    paid INTEGER DEFAULT 0,
    paid_amount REAL DEFAULT 0,
    koc_paid_amount REAL DEFAULT 0,
    hop_dong_mua TEXT,
    unc_mua TEXT,
    hop_dong_ban TEXT,
    unc_ban TEXT,
    note TEXT
)
""")


cursor.execute("""
CREATE TABLE IF NOT EXISTS total_revenue_targets (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    target_month TEXT UNIQUE,
    target_revenue REAL DEFAULT 0,
    note TEXT
)
""")


cursor.execute("""
CREATE TABLE IF NOT EXISTS app_settings (
    key TEXT PRIMARY KEY,
    value BLOB
)
""")

conn.commit()

def get_setting(key):
    row = cursor.execute("SELECT value FROM app_settings WHERE key = ?", (key,)).fetchone()
    return row[0] if row else None

def save_setting(key, value):
    cursor.execute("""INSERT INTO app_settings(key, value) VALUES (?, ?) ON CONFLICT(key) DO UPDATE SET value = excluded.value""", (key, value))
    conn.commit()

def delete_setting(key):
    cursor.execute("DELETE FROM app_settings WHERE key = ?", (key,))
    conn.commit()

def image_data_uri(data, mime):
    if not data:
        return None
    return f"data:{mime};base64,{base64.b64encode(data).decode('utf-8')}"


# =========================================================
# DATABASE MIGRATION
# =========================================================
# Nếu bảng booking_services đã tồn tại từ bản cũ thì tự thêm
# cột tiền KOC đã thanh toán, không làm mất dữ liệu cũ.
# =========================================================

booking_columns = [
    row[1]
    for row in cursor.execute(
        "PRAGMA table_info(booking_services)"
    ).fetchall()
]

booking_migrations = {
    "koc_paid_amount": "REAL DEFAULT 0",
    "hop_dong_mua": "TEXT",
    "unc_mua": "TEXT",
    "hop_dong_ban": "TEXT",
    "unc_ban": "TEXT",
}

for col_name, col_type in booking_migrations.items():
    if col_name not in booking_columns:
        cursor.execute(
            f"ALTER TABLE booking_services ADD COLUMN {col_name} {col_type}"
        )

conn.commit()


# =========================================================

# =========================================================
# TẠO HỢP ĐỒNG WORD MỚI
# =========================================================
def number_to_vietnamese(n):
    n=int(round(float(n or 0)))
    if n==0: return "Không đồng"
    d=["không","một","hai","ba","bốn","năm","sáu","bảy","tám","chín"]
    def r3(x, full=False):
        a,b,c=x//100,(x//10)%10,x%10; s=[]
        if a or full: s += [d[a],"trăm"]
        if b==0 and c and (a or full): s += ["lẻ"]
        elif b: s += [d[b],"mươi"]
        if c:
            if b>=2 and c==1: w="mốt"
            elif b>=1 and c==5: w="lăm"
            elif b>=2 and c==4: w="tư"
            else: w=d[c]
            s.append(w)
        return " ".join(s)
    units=["","nghìn","triệu","tỷ"]; groups=[]
    while n: groups.append(n%1000); n//=1000
    out=[]
    for i in range(len(groups)-1,-1,-1):
        if groups[i]: out += [r3(groups[i], i==len(groups)-1), units[i]] if i else [r3(groups[i], True)]
    return " ".join(out).capitalize()+" đồng"

def make_booking_doc(data, path):
    from docx import Document
    from docx.shared import Cm, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    doc=Document(); sec=doc.sections[0]
    sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Cm(2.2)
    doc.styles["Normal"].font.name="Times New Roman"; doc.styles["Normal"].font.size=Pt(12)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\nĐộc lập – Tự do – Hạnh phúc"); r.bold=True
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run("HỢP ĐỒNG CUNG CẤP KOLs/KOCs"); r.bold=True; r.font.size=Pt(14)
    p=doc.add_paragraph(f"Số: {data['number']}-{data['year']}/SINGO-KOL"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(f"Hôm nay, ngày {data['day']} tháng {data['month']} năm {data['year']}, chúng tôi gồm có:")
    for x in ["BÊN A: CÔNG TY TNHH SINGO","Người đại diện: Ông NGUYỄN THÁI HỢP – Chức vụ: Giám Đốc","Địa chỉ: 125/48/13 Lê Đức Thọ, Phường Gò Vấp, Thành phố Hồ Chí Minh, Việt Nam.","Mã số thuế: 0314437471",f"BÊN B: {data['partner']}",f"Người đại diện: {data['rep']}",f"Chức vụ: {data['position']}",f"Địa chỉ: {data['address']}",f"Số điện thoại: {data['phone']}",f"Mã số thuế: {data['tax']}"]: doc.add_paragraph(x)
    sections=[
      ("ĐIỀU 1. NỘI DUNG CÔNG VIỆC",[f"Bên A cung cấp dịch vụ KOL/KOC theo gói {data['package']} với số lượng {data['qty']} gói trên nền tảng {data['platform']}.",f"KOC không phí: {data['free']}; KOL Micro có phí: {data['micro']}; KOC/KOL Macro có phí: {data['macro']}.",f"Hoa hồng tiêu chuẩn: {data['commission']}. Hoa hồng Ads: {data['ads']}.",f"Tiêu chí lựa chọn/yêu cầu chiến dịch: {data['criteria']}."]),
      ("ĐIỀU 2. GIÁ TRỊ HỢP ĐỒNG",[f"Phí quản lý Agency: {data['fee']:,.0f} VNĐ.",f"Thuế VAT ({data['vat_rate']:.0f}%): {data['vat']:,.0f} VNĐ.",f"Tổng giá trị Hợp đồng: {data['total']:,.0f} VNĐ.",f"Bằng chữ: {number_to_vietnamese(data['total'])}.","Phí booking KOL/KOC không nằm trong Giá trị Hợp đồng và chỉ phát sinh khi Hai Bên có thỏa thuận bằng văn bản hoặc phụ lục."]),
      ("ĐIỀU 3. THANH TOÁN",["Bên B thanh toán 100% Giá trị Hợp đồng bằng chuyển khoản theo thông tin do Bên A cung cấp sau khi Hai Bên ký kết, trừ khi có thỏa thuận khác bằng văn bản."]),
      ("ĐIỀU 4. QUYỀN VÀ NGHĨA VỤ",["Bên A thực hiện dịch vụ đúng phạm vi công việc; Bên B cung cấp thông tin, tài liệu, sản phẩm cần thiết và thanh toán đúng hạn.","Các Bên phối hợp xác nhận nội dung, danh sách KOL/KOC và tiến độ triển khai theo từng chiến dịch."]),
      ("ĐIỀU 5. BẢO MẬT VÀ BẢN QUYỀN",["Các Bên bảo mật thông tin nhận được trong quá trình hợp tác. Quyền sử dụng nội dung được thực hiện theo thỏa thuận chiến dịch, email hoặc phụ lục được xác nhận bởi Hai Bên."]),
      ("ĐIỀU 6. HIỆU LỰC VÀ GIẢI QUYẾT TRANH CHẤP",["Hợp đồng có hiệu lực từ ngày ký đến khi hoàn tất nghĩa vụ. Mọi sửa đổi phải lập thành văn bản. Tranh chấp được ưu tiên thương lượng và giải quyết theo pháp luật Việt Nam."])]
    for title,paras in sections:
        p=doc.add_paragraph(); p.add_run(title).bold=True
        for x in paras: doc.add_paragraph(x)
    t=doc.add_table(rows=1,cols=2); t.style="Table Grid"; t.cell(0,0).text="ĐẠI DIỆN BÊN A\nCÔNG TY TNHH SINGO\n\n\nNGUYỄN THÁI HỢP"; t.cell(0,1).text=f"ĐẠI DIỆN BÊN B\n{data['partner']}\n\n\n{data['rep']}"
    doc.save(path)

def make_ctv_doc(data,path):
    from docx import Document
    from docx.shared import Cm, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    doc=Document(); sec=doc.sections[0]; sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Cm(2.2)
    doc.styles["Normal"].font.name="Times New Roman"; doc.styles["Normal"].font.size=Pt(12)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\nĐộc lập – Tự do – Hạnh phúc"); r.bold=True
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run("HỢP ĐỒNG CỘNG TÁC VIÊN"); r.bold=True; r.font.size=Pt(14)
    p=doc.add_paragraph(f"Số: {data['number']}/{data['year']}/HĐCTV - GP/37471 - 001"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(f"Hợp đồng được lập ngày {data['day']} tháng {data['month']} năm {data['year']} giữa các bên:")
    for x in ["BÊN A: CÔNG TY TNHH SINGO","Địa chỉ: 125/48/13 Lê Đức Thọ, Phường Gò Vấp, Thành phố Hồ Chí Minh, Việt Nam.","Mã số thuế: 0314437471","Đại diện: Ông NGUYỄN THÁI HỢP – Chức vụ: Giám Đốc",f"BÊN B: {data['name']}",f"Sinh ngày: {data['dob']}",f"CCCD: {data['cccd']} – Ngày cấp: {data['cccd_date']} – Nơi cấp: {data['cccd_place']}",f"Mã số thuế: {data['tax']}",f"Số tài khoản: {data['account']} – Ngân hàng: {data['bank']}"]: doc.add_paragraph(x)
    for title,text in [("ĐIỀU 1. ĐỐI TƯỢNG HỢP ĐỒNG",f"Bên B thực hiện công việc cho nhãn hàng {data['brand']} trên nền tảng TikTok."),("ĐIỀU 2. NỘI DUNG DỊCH VỤ",f"KOL/KOC: {data['username']}. Hạng mục: {data['task']}. Số lượng: {data['qty']}.")]:
        p=doc.add_paragraph(); p.add_run(title).bold=True; doc.add_paragraph(text)
    t=doc.add_table(rows=2,cols=4); t.style="Table Grid"
    for i,h in enumerate(["STT","Hạng mục","SL","Thành tiền (VNĐ)"]): t.cell(0,i).text=h
    t.cell(1,0).text="1"; t.cell(1,1).text=data['task']; t.cell(1,2).text=str(data['qty']); t.cell(1,3).text=f"{data['gross']:,.0f}"
    p=doc.add_paragraph(); p.add_run("ĐIỀU 3. THÙ LAO VÀ THANH TOÁN").bold=True
    doc.add_paragraph(f"Tổng giá trị thù lao: {data['gross']:,.0f} VNĐ.\nThuế TNCN khấu trừ ({data['tax_rate']:.0f}%): {data['tax']:,.0f} VNĐ.\nCTV thực nhận: {data['net']:,.0f} VNĐ.\nBằng chữ: {number_to_vietnamese(data['net'])}.")
    doc.add_paragraph("Bên A khấu trừ thuế TNCN theo tỷ lệ được thỏa thuận tại Hợp đồng. Bên A thanh toán bằng chuyển khoản trong vòng 15 ngày làm việc kể từ khi nhận Hợp đồng đã ký và xác nhận hoàn thành công việc.")
    for title,text in [("ĐIỀU 4. TIÊU CHUẨN DỊCH VỤ","Bên B thực hiện đúng yêu cầu, tiến độ và chịu trách nhiệm chỉnh sửa sai sót thuộc phạm vi công việc."),("ĐIỀU 5. QUYỀN VÀ NGHĨA VỤ","Các Bên phối hợp thực hiện, nghiệm thu, thanh toán và bảo mật thông tin theo Hợp đồng."),("ĐIỀU 6. BẢN QUYỀN VÀ BẢO MẬT","Nội dung và thông tin phát sinh trong phạm vi Hợp đồng được sử dụng theo thỏa thuận của Các Bên; nghĩa vụ bảo mật tiếp tục có hiệu lực sau khi Hợp đồng kết thúc."),("ĐIỀU 7. HIỆU LỰC VÀ TRANH CHẤP","Hợp đồng có hiệu lực từ ngày ký đến khi hoàn tất nghĩa vụ. Mọi sửa đổi phải lập thành văn bản. Tranh chấp được ưu tiên thương lượng và giải quyết theo pháp luật Việt Nam.")]:
        p=doc.add_paragraph(); p.add_run(title).bold=True; doc.add_paragraph(text)
    t=doc.add_table(rows=1,cols=2); t.style="Table Grid"; t.cell(0,0).text="ĐẠI DIỆN BÊN A\nCÔNG TY TNHH SINGO\n\n\nNGUYỄN THÁI HỢP"; t.cell(0,1).text=f"ĐẠI DIỆN BÊN B\n{data['name']}\n\n\n{data['name']}"
    doc.save(path)

# PAGE CONFIG
# =========================================================


# =========================
# QUẢN LÝ HỢP ĐỒNG
# =========================
def ensure_contract_tables():
    conn = sqlite3.connect(DB_PATH)
    cur = conn.cursor()
    cur.execute("""
        CREATE TABLE IF NOT EXISTS contracts (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            contract_type TEXT NOT NULL,
            contract_number INTEGER NOT NULL,
            contract_year INTEGER NOT NULL,
            contract_code TEXT NOT NULL,
            signed_date TEXT,
            partner_name TEXT,
            brand_name TEXT,
            contract_value REAL DEFAULT 0,
            status TEXT DEFAULT 'Dự thảo',
            note TEXT,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP
        )
    """)
    cur.execute("""
        CREATE TABLE IF NOT EXISTS contract_settings (
            contract_type TEXT PRIMARY KEY,
            start_number INTEGER NOT NULL
        )
    """)
    cur.execute(
        "INSERT OR IGNORE INTO contract_settings(contract_type,start_number) VALUES (?,?)",
        ("Booking", 75)
    )
    cur.execute(
        "INSERT OR IGNORE INTO contract_settings(contract_type,start_number) VALUES (?,?)",
        ("CTV", 140)
    )
    conn.commit()
    conn.close()

def get_next_contract_number(contract_type):
    conn = sqlite3.connect(DB_PATH)
    cur = conn.cursor()
    cur.execute("SELECT start_number FROM contract_settings WHERE contract_type=?", (contract_type,))
    row = cur.fetchone()
    start = int(row[0]) if row else (75 if contract_type == "Booking" else 140)
    cur.execute(
        "SELECT contract_number FROM contracts WHERE contract_type=? ORDER BY contract_number",
        (contract_type,)
    )
    used = {int(r[0]) for r in cur.fetchall()}
    n = start
    while n in used:
        n += 1
    conn.close()
    return n

def save_contract(contract_type, contract_number, year, partner_name="", brand_name="",
                  signed_date="", contract_value=0, status="Dự thảo", note=""):
    conn = sqlite3.connect(DB_PATH)
    conn.execute("""
        INSERT INTO contracts
        (contract_type, contract_number, contract_year, contract_code, signed_date,
         partner_name, brand_name, contract_value, status, note)
        VALUES (?,?,?,?,?,?,?,?,?,?)
    """, (
        contract_type, int(contract_number), int(year),
        f"{int(contract_number)}-{int(year)}", signed_date,
        partner_name, brand_name, float(contract_value or 0), status, note
    ))
    conn.commit()
    conn.close()

def delete_contract(contract_id):
    conn = sqlite3.connect(DB_PATH)
    conn.execute("DELETE FROM contracts WHERE id=?", (int(contract_id),))
    conn.commit()
    conn.close()

def get_contracts():
    conn = sqlite3.connect(DB_PATH)
    df = pd.read_sql_query(
        "SELECT * FROM contracts ORDER BY contract_year DESC, contract_number DESC",
        conn
    )
    conn.close()
    return df

ensure_contract_tables()

st.set_page_config(
    page_title="SINGO AGENCY | KOC Management & Analytics",
    page_icon="✦",
    layout="wide",
    initial_sidebar_state="expanded"
)


# =========================================================
# FUNCTIONS
# =========================================================

def money(value):
    return f"{value:,.0f} đ"


def percent(value):
    return f"{value:.2f}%"


def normalize_month_value(value):
    """Chuẩn hóa tháng về YYYY-MM để các nguồn dữ liệu match nhau."""
    if pd.isna(value):
        return ""

    text = str(value).strip()

    if len(text) == 7 and text[4] == "-":
        return text

    parsed = pd.to_datetime(value, errors="coerce")

    if pd.isna(parsed):
        return text

    return parsed.strftime("%Y-%m")


def get_target_cycle(month_text):
    """
    Công ty tính Target theo chu kỳ 2 tháng:
    09-10, 11-12, 01-02, 03-04, ...
    """
    year, month = map(int, str(month_text).split("-"))

    if month % 2 == 1:
        start_month = month
        end_month = month + 1
    else:
        start_month = month - 1
        end_month = month

    # Xử lý trường hợp tháng 12.
    if end_month == 13:
        start_month = 11
        end_month = 12

    cycle_start = f"{year:04d}-{start_month:02d}"
    cycle_end = f"{year:04d}-{end_month:02d}"

    return cycle_start, cycle_end, f"{start_month:02d}/{year} - {end_month:02d}/{year}"


def clean_money(series):
    """
    Xử lý tiền TikTok dạng:
    17.683.683.391₫
    708.471.259₫
    0₫
    """

    return pd.to_numeric(
        series
        .astype(str)
        .str.replace("₫", "", regex=False)
        .str.replace("đ", "", regex=False)
        .str.replace(",", "", regex=False)
        .str.replace(".", "", regex=False)
        .str.replace(" ", "", regex=False)
        .str.strip(),
        errors="coerce"
    ).fillna(0)


def parse_tiktok_date(series):
    """
    File TikTok có dạng:

    2026-09-01-2026-09-30

    Lấy ngày đầu tiên:

    2026-09-01
    """

    text = (
        series
        .astype(str)
        .str.strip()
    )

    extracted = text.str.extract(
        r"(\d{4}-\d{2}-\d{2})",
        expand=False
    )

    return pd.to_datetime(
        extracted,
        format="%Y-%m-%d",
        errors="coerce"
    )


# =========================================================
# GIAO DIỆN / THEME
# =========================================================

logo_bytes = get_setting("logo")
bg_bytes = get_setting("background")
logo_mime = get_setting("logo_mime") or "image/png"
bg_mime = get_setting("background_mime") or "image/jpeg"
logo_uri = image_data_uri(logo_bytes, logo_mime)
bg_uri = image_data_uri(bg_bytes, bg_mime)
bg_css = f"url('{bg_uri}')" if bg_uri else "linear-gradient(135deg, #fff8fa 0%, #fdecef 45%, #f8e1e7 100%)"

st.markdown(
    f"""
    <style>
    .stApp {{ background: {bg_css} center center / cover fixed no-repeat; }}
    .stApp::before {{ content: ""; position: fixed; inset: 0; background: rgba(255,248,250,.82); z-index:-1; }}
    [data-testid="stSidebar"] {{ background: linear-gradient(180deg,rgba(255,255,255,.98),rgba(255,244,247,.97)); border-right:1px solid #f0d6dd; }}
    .app-logo-wrap {{ display:flex; align-items:center; gap:12px; padding:8px 4px 18px; border-bottom:1px solid #f2dbe1; margin-bottom:16px; }}
    .app-logo {{ width:54px; height:54px; object-fit:contain; border-radius:14px; background:white; border:1px solid #f0d6dd; padding:5px; }}
    .app-logo-title {{ font-size:18px; font-weight:850; letter-spacing:.02em; line-height:1.05; color:#7e3145; }}
    .app-logo-sub {{ font-size:11px; color:#8b777c; margin-top:4px; }}
    .menu-label {{ font-size:11px; text-transform:uppercase; letter-spacing:.08em; color:#a17984; font-weight:700; margin:8px 0 6px 4px; }}
    div.stButton > button {{ border-radius:14px; border:1px solid transparent; background:rgba(255,255,255,.42); color:#4b343a; font-weight:600; min-height:44px; box-shadow:none; transition:all .18s ease; }}
    div.stButton > button:hover {{ border-color:#edc1cc; color:#a83f58; background:#fff2f5; transform:translateX(2px); }}
    .nav-active button {{ background:linear-gradient(90deg,#f8dbe3,#fdeef2) !important; border-color:#efc2cd !important; color:#9e3b55 !important; font-weight:800 !important; box-shadow:0 4px 12px rgba(170,75,100,.08) !important; }}
    .nav-child button {{ text-align:left !important; padding-left:16px !important; font-size:13px !important; background:transparent !important; border-color:transparent !important; min-height:38px !important; }}
    [data-testid="stMetric"] {{ background:rgba(255,255,255,.88); border:1px solid #f0d9df; padding:12px 13px; border-radius:16px; box-shadow:0 5px 18px rgba(150,75,95,.06); min-width:0; }}
    [data-testid="stMetricLabel"] {{ font-size:0.78rem !important; line-height:1.15 !important; white-space:nowrap; overflow:hidden; text-overflow:ellipsis; }}
    [data-testid="stMetricValue"] {{ font-size:22px !important; line-height:1.15 !important; white-space:nowrap !important; overflow:visible !important; text-overflow:clip !important; letter-spacing:-0.4px; }}
    [data-testid="stMetricValue"] > div {{ white-space:nowrap !important; overflow:visible !important; }}
    h1,h2,h3 {{ color:#3a252b; }}
    </style>
    """, unsafe_allow_html=True
)

if logo_uri:
    logo_html = f'<img class="app-logo" src="{logo_uri}">'
else:
    logo_html = '<div class="app-logo" style="display:flex;align-items:center;justify-content:center;font-size:25px;">✦</div>'

with st.sidebar:
    st.markdown(
        f"<div class=\"app-logo-wrap\">{logo_html}<div><div class=\"app-logo-title\">SINGO AGENCY</div><div class=\"app-logo-sub\">KOC Management & Analytics</div></div></div>",
        unsafe_allow_html=True
    )
    if "page" not in st.session_state:
        st.session_state.page = "dashboard"
    if "dash_expanded" not in st.session_state:
        st.session_state.dash_expanded = True

    st.markdown('<div class="menu-label">MENU</div>', unsafe_allow_html=True)

    # Menu cha: Tổng quan TAP / Booking
    dash_label = "⌄  Tổng quan số liệu TAP/Booking" if st.session_state.dash_expanded else "›  Tổng quan số liệu TAP/Booking"
    dash_active = st.session_state.page in ["dashboard", "booking", "tap_target", "monthly"]
    st.markdown(f'<div class="{"nav-active" if dash_active and st.session_state.page == "dashboard" else ""}">', unsafe_allow_html=True)
    if st.button(dash_label, use_container_width=True, key="nav_dash"):
        st.session_state.page = "dashboard"
        st.session_state.dash_expanded = not st.session_state.dash_expanded
        st.rerun()
    st.markdown('</div>', unsafe_allow_html=True)

    if st.session_state.dash_expanded:
        for label,key in [("▱  Dịch vụ Booking","booking"),("◎  Phân tích TAP","tap_target"),("▥  MO/DA (Monthly Analytics)","monthly")]:
            active = st.session_state.page == key
            st.markdown(f'<div class="{"nav-active " if active else ""}nav-child">', unsafe_allow_html=True)
            if st.button(label, use_container_width=True, key=f"nav_{key}"):
                st.session_state.page = key
                st.rerun()
            st.markdown('</div>', unsafe_allow_html=True)

    st.markdown('<div style="height:10px"></div>', unsafe_allow_html=True)

    for label,key in [("♙  DATA KOC lưu trữ","koc"),("◇  Danh sách Brand","brand"),("📄  Hợp đồng","contracts"),("⚙  Cài đặt","settings")]:
        active = st.session_state.page == key
        st.markdown(f'<div class="{"nav-active" if active else ""}">', unsafe_allow_html=True)
        if st.button(label, use_container_width=True, key=f"nav_{key}"):
            st.session_state.page = key
            st.rerun()
        st.markdown('</div>', unsafe_allow_html=True)

page = st.session_state.page
page_titles = {"dashboard":"Tổng quan số liệu TAP/Booking","booking":"Dịch vụ Booking","tap_target":"Phân tích TAP","monthly":"MO/DA (Monthly Analytics)","koc":"DATA KOC lưu trữ","brand":"Danh sách Brand","contracts":"Hợp đồng","settings":"Cài đặt"}
st.title(page_titles.get(page,"KOC Management & Analytics"))
if page != "settings":
    st.caption("Hệ thống quản lý KOC/KOL, Brand, TAP và dịch vụ Booking")


# =========================================================
# DASHBOARD
# =========================================================



# =========================
# TRANG HỢP ĐỒNG
if page == "contracts":
    st.header("📄 Quản lý Hợp đồng")
    st.caption("Tạo hợp đồng Word mới. Nút xóa hợp đồng vẫn được giữ.")
    tab_create,tab_list,tab_settings=st.tabs(["➕ Tạo hợp đồng","📋 Danh sách","⚙️ Số hợp đồng"])
    with tab_create:
        contract_type=st.selectbox("Loại hợp đồng",["Booking","CTV"])
        next_no=get_next_contract_number(contract_type)
        st.info(f"Số hợp đồng đề xuất tiếp theo: **{next_no}**")
        c1,c2=st.columns(2)
        with c1:
            contract_number=st.number_input("Số hợp đồng",min_value=1,value=next_no,step=1)
            year=st.number_input("Năm",min_value=2020,max_value=2100,value=2026,step=1)
            signed_date=st.date_input("Ngày ký")
        with c2: status=st.selectbox("Trạng thái",["Dự thảo","Đã ký","Đã thanh toán","Đã hủy"])
        if contract_type=="Booking":
            st.subheader("Thông tin Bên B")
            c1,c2=st.columns(2)
            with c1:
                partner=st.text_input("Tên công ty / đối tác"); rep=st.text_input("Người đại diện"); position=st.text_input("Chức vụ",value="Giám Đốc"); phone=st.text_input("Số điện thoại")
            with c2:
                address=st.text_area("Địa chỉ"); tax=st.text_input("Mã số thuế")
            st.subheader("Nội dung chiến dịch")
            c1,c2,c3=st.columns(3)
            with c1: package=st.text_input("Tên gói",value="B1"); package_qty=st.number_input("Số lượng gói",min_value=1,value=1,step=1); platform=st.text_input("Nền tảng",value="TikTok")
            with c2: free=st.number_input("KOC không phí",min_value=0,value=0); micro=st.number_input("KOL Micro có phí",min_value=0,value=0); macro=st.number_input("KOC/KOL Macro có phí",min_value=0,value=0)
            with c3: commission=st.text_input("Hoa hồng tiêu chuẩn"); ads=st.text_input("Hoa hồng Ads"); fee=st.number_input("Phí quản lý Agency (VNĐ)",min_value=0.0,value=0.0,step=100000.0)
            criteria=st.text_area("Tiêu chí / yêu cầu chiến dịch"); vat_rate=st.number_input("VAT (%)",min_value=0.0,value=8.0,step=0.5)
            vat=fee*vat_rate/100; total=fee+vat; st.metric("Tổng giá trị hợp đồng",f"{total:,.0f} VNĐ"); note=st.text_area("Ghi chú")
            if st.button("📄 Soạn & tải Hợp đồng Word",type="primary",use_container_width=True):
                if not partner.strip(): st.error("Vui lòng nhập tên đối tác.")
                else:
                    con=sqlite3.connect(DB_PATH); exists=con.execute("SELECT 1 FROM contracts WHERE contract_type=? AND contract_number=? AND contract_year=?",("Booking",int(contract_number),int(year))).fetchone(); con.close()
                    if exists: st.error(f"Số {contract_number} đã tồn tại.")
                    else:
                        path=f"/mnt/data/HĐ_Booking_{int(contract_number)}_{int(year)}.docx"
                        make_booking_doc({'number':int(contract_number),'year':int(year),'day':signed_date.day,'month':signed_date.month,'partner':partner,'rep':rep,'position':position,'phone':phone,'address':address,'tax':tax,'package':package,'qty':package_qty,'platform':platform,'free':free,'micro':micro,'macro':macro,'commission':commission,'ads':ads,'fee':fee,'vat_rate':vat_rate,'vat':vat,'total':total,'criteria':criteria},path)
                        save_contract("Booking",contract_number,year,partner,"",str(signed_date),total,status,note)
                        st.success("Đã soạn và lưu HĐ Booking.")
                        with open(path,"rb") as f: st.download_button("⬇️ Tải HĐ Booking Word",f,file_name=Path(path).name,mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",use_container_width=True)
        else:
            st.subheader("Thông tin CTV")
            c1,c2=st.columns(2)
            with c1: name=st.text_input("Họ và tên CTV"); dob=st.text_input("Ngày sinh"); cccd=st.text_input("CCCD"); cccd_date=st.text_input("Ngày cấp CCCD"); cccd_place=st.text_input("Nơi cấp CCCD")
            with c2: tax=st.text_input("Mã số thuế"); account=st.text_input("Số tài khoản"); bank=st.text_input("Ngân hàng"); brand=st.text_input("Nhãn hàng"); username=st.text_input("KOL/KOC / Username")
            task=st.text_area("Hạng mục công việc",value="Tham gia sản xuất và đăng tải video trên kênh TikTok với nội dung giới thiệu sản phẩm theo định hướng nhãn hàng.")
            c1,c2,c3=st.columns(3)
            with c1: qty=st.number_input("Số lượng",min_value=0.0,value=1.0,step=1.0)
            with c2: unit=st.number_input("Đơn giá (VNĐ)",min_value=0.0,value=0.0,step=10000.0)
            with c3: tax_rate=st.number_input("Thuế TNCN (%)",min_value=0.0,value=10.0,step=0.5)
            gross=qty*unit; tax_amt=gross*tax_rate/100; net=gross-tax_amt; c1,c2=st.columns(2); c1.metric("Tổng thù lao",f"{gross:,.0f} VNĐ"); c2.metric("CTV thực nhận",f"{net:,.0f} VNĐ"); note=st.text_area("Ghi chú")
            if st.button("📄 Soạn & tải Hợp đồng Word",type="primary",use_container_width=True):
                if not name.strip(): st.error("Vui lòng nhập họ tên CTV.")
                else:
                    con=sqlite3.connect(DB_PATH); exists=con.execute("SELECT 1 FROM contracts WHERE contract_type=? AND contract_number=? AND contract_year=?",("CTV",int(contract_number),int(year))).fetchone(); con.close()
                    if exists: st.error(f"Số {contract_number} đã tồn tại.")
                    else:
                        path=f"/mnt/data/HĐ_CTV_{int(contract_number)}_{int(year)}.docx"
                        make_ctv_doc({'number':int(contract_number),'year':int(year),'day':signed_date.day,'month':signed_date.month,'name':name,'dob':dob,'cccd':cccd,'cccd_date':cccd_date,'cccd_place':cccd_place,'tax':tax,'account':account,'bank':bank,'brand':brand,'username':username,'task':task,'qty':qty,'gross':gross,'tax_rate':tax_rate,'tax':tax_amt,'net':net},path)
                        save_contract("CTV",contract_number,year,name,brand,str(signed_date),net,status,note)
                        st.success("Đã soạn và lưu HĐ CTV.")
                        with open(path,"rb") as f: st.download_button("⬇️ Tải HĐ CTV Word",f,file_name=Path(path).name,mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",use_container_width=True)
    with tab_list:
        df_contracts=get_contracts()
        if df_contracts.empty: st.info("Chưa có hợp đồng nào.")
        else:
            display=df_contracts[["id","contract_type","contract_number","contract_year","signed_date","partner_name","brand_name","contract_value","status"]].copy(); display.columns=["ID","Loại HĐ","Số HĐ","Năm","Ngày ký","Đối tác/CTV","Brand","Giá trị","Trạng thái"]; display["Giá trị"]=display["Giá trị"].map(lambda x:f"{x:,.0f} đ"); st.dataframe(display,use_container_width=True,hide_index=True)
            st.markdown("### 🗑️ Xóa hợp đồng")
            del_id=st.number_input("ID hợp đồng cần xóa",min_value=1,step=1)
            if st.button("🗑️ Xóa hợp đồng",type="secondary"): delete_contract(del_id); st.success("Đã xóa. Số hợp đồng được trả lại để cấp lại."); st.rerun()
    with tab_settings:
        st.write("### Số bắt đầu")
        con=sqlite3.connect(DB_PATH); rows=con.execute("SELECT contract_type,start_number FROM contract_settings ORDER BY contract_type").fetchall(); con.close()
        for ctype,start in rows:
            new_start=st.number_input(f"HĐ {ctype} - số bắt đầu",min_value=1,value=int(start),step=1,key=f"start_{ctype}")
            if st.button(f"Lưu số bắt đầu {ctype}",key=f"save_start_{ctype}"): con=sqlite3.connect(DB_PATH); con.execute("UPDATE contract_settings SET start_number=? WHERE contract_type=?",(int(new_start),ctype)); con.commit(); con.close(); st.rerun()
    st.stop()

# DASHBOARD DATA
    # =====================================================

    analytics_data = pd.read_sql(
        """
        SELECT *
        FROM analytics
        """,
        conn
    )


    if analytics_data.empty:

        st.warning(
            "Chưa có dữ liệu Analytics. "
            "Hãy upload file TikTok ở trên."
        )

    else:

        st.divider()

        st.subheader(
            "📊 Dashboard"
        )


        # =================================================
        # FILTER MONTH
        # =================================================

        months = sorted(
            analytics_data[
                "report_month"
            ].dropna().unique(),
            reverse=True
        )


        col1, col2 = st.columns(2)


        with col1:

            selected_month = st.selectbox(
                "📅 Chọn tháng",
                months
            )


        month_data = analytics_data[
            analytics_data["report_month"]
            == selected_month
        ].copy()


        # =================================================
        # FILTER BRAND
        # =================================================

        brands = [
            "Tất cả"
        ] + sorted(
            month_data[
                "store_name"
            ]
            .dropna()
            .unique()
            .tolist()
        )


        with col2:

            selected_brand = st.selectbox(
                "🏷️ Chọn Brand",
                brands
            )


        if selected_brand != "Tất cả":

            month_data = month_data[
                month_data["store_name"]
                == selected_brand
            ]


        # =================================================
        # KPI
        # =================================================

        total_gmv = (
            month_data["gmv"]
            .sum()
        )


        total_ac = (
            month_data["ac"]
            .sum()
        )


        ac_rate = (
            total_ac
            / total_gmv
            * 100
            if total_gmv
            else 0
        )


        col1, col2, col3 = st.columns(3)


        with col1:

            st.metric(
                "💰 Tổng GMV",
                money(total_gmv)
            )


        with col2:

            st.metric(
                "💵 Tổng hoa hồng thực tế",
                money(total_ac)
            )


        with col3:

            st.metric(
                "📊 Tỷ lệ hoa hồng",
                percent(ac_rate)
            )


        st.divider()


        # =================================================
        # TOP BRAND
        # =================================================

        st.subheader(
            "🏆 Top Brand"
        )


        top_brand = (
            month_data
            .groupby(
                "store_name",
                as_index=False
            )
            .agg(
                GMV=("gmv", "sum"),
                AC=("ac", "sum")
            )
            .sort_values(
                "AC",
                ascending=False
            )
            .head(20)
        )


        top_brand["Tỷ lệ hoa hồng"] = (
            top_brand["AC"]
            / top_brand["GMV"]
            * 100
        ).fillna(0)


        brand_display = top_brand.copy()


        brand_display["GMV"] = (
            brand_display["GMV"]
            .apply(money)
        )


        brand_display["Hoa hồng thực tế"] = (
            brand_display["AC"]
            .apply(money)
        )


        st.dataframe(
            brand_display,
            use_container_width=True,
            hide_index=True
        )


        # =================================================
        # TOP KOC
        # =================================================

        st.subheader(
            "👑 Top KOC/NST"
        )


        top_koc = (
            month_data
            .groupby(
                "creator_name",
                as_index=False
            )
            .agg(
                GMV=("gmv", "sum"),
                AC=("ac", "sum")
            )
            .sort_values(
                "AC",
                ascending=False
            )
            .head(20)
        )


        top_koc["Tỷ lệ hoa hồng"] = (
            top_koc["AC"]
            / top_koc["GMV"]
            * 100
        ).fillna(0)


        koc_display = top_koc.copy()


        koc_display["GMV"] = (
            koc_display["GMV"]
            .apply(money)
        )


        koc_display["Hoa hồng thực tế"] = (
            koc_display["AC"]
            .apply(money)
        )


        st.dataframe(
            koc_display,
            use_container_width=True,
            hide_index=True
        )


        # =================================================
        # BRAND ANALYSIS
        # =================================================

        st.divider()

        st.subheader(
            "📊 Phân tích Brand"
        )


        brand_analysis = (
            month_data
            .groupby(
                "store_name",
                as_index=False
            )
            .agg(
                GMV=("gmv", "sum"),
                AC=("ac", "sum"),
                KOC=("creator_name", "nunique")
            )
        )


        brand_analysis["Tỷ lệ hoa hồng"] = (
            brand_analysis["AC"]
            / brand_analysis["GMV"]
            * 100
        ).fillna(0)


        brand_analysis = (
            brand_analysis
            .sort_values(
                "AC",
                ascending=False
            )
        )


        analysis_display = (
            brand_analysis.copy()
        )


        analysis_display["GMV"] = (
            analysis_display["GMV"]
            .apply(money)
        )


        analysis_display["Hoa hồng thực tế"] = (
            analysis_display["AC"]
            .apply(money)
        )


        st.dataframe(
            analysis_display,
            use_container_width=True,
            hide_index=True
        )


        # =================================================
        # MONTHLY SUMMARY
        # =================================================

        st.divider()

        st.subheader(
            "📅 Tổng quan theo tháng"
        )


        monthly = (
            analytics_data
            .groupby(
                "report_month",
                as_index=False
            )
            .agg(
                GMV=("gmv", "sum"),
                AC=("ac", "sum"),
                KOC=("creator_name", "nunique"),
                Brand=("store_name", "nunique")
            )
        )


        monthly["Tỷ lệ hoa hồng"] = (
            monthly["AC"]
            / monthly["GMV"]
            * 100
        ).fillna(0)


        monthly_display = monthly.copy()


        monthly_display["GMV"] = (
            monthly_display["GMV"]
            .apply(money)
        )


        monthly_display["Hoa hồng thực tế"] = (
            monthly_display["AC"]
            .apply(money)
        )


        st.dataframe(
            monthly_display,
            use_container_width=True,
            hide_index=True
        )


        # =================================================
        # DETAIL
        # =================================================

        with st.expander(
            "🔍 Xem dữ liệu chi tiết"
        ):

            detail = month_data[
                [
                    "report_date",
                    "creator_name",
                    "store_name",
                    "gmv",
                    "ac"
                ]
            ].copy()


            detail["report_date"] = (
                pd.to_datetime(
                    detail["report_date"],
                    errors="coerce"
                )
                .dt.strftime("%d/%m/%Y")
            )


            detail["gmv"] = (
                detail["gmv"]
                .apply(money)
            )


            detail["ac"] = (
                detail["ac"]
                .apply(money)
            )


            detail.columns = [
                "Ngày",
                "KOC/NST",
                "Brand",
                "GMV",
                "Hoa hồng thực tế"
            ]


            st.dataframe(
                detail,
                use_container_width=True,
                hide_index=True
            )


        # =================================================
        # DOWNLOAD
        # =================================================

        csv = (
            month_data
            .to_csv(
                index=False
            )
            .encode("utf-8-sig")
        )


        st.download_button(
            "📥 Xuất dữ liệu tháng",
            data=csv,
            file_name=(
                f"analytics_{selected_month}.csv"
            ),
            mime="text/csv"
        )


elif page == "settings":
    st.header("Cài đặt giao diện")
    st.caption("Tuỳ chỉnh logo, hình nền và giao diện hồng pastel cho toàn bộ web.")
    tab1, tab2 = st.tabs(["🎨 Giao diện", "👀 Xem trước"])
    with tab1:
        st.subheader("Logo website")
        logo_upload = st.file_uploader("Tải logo", type=["png","jpg","jpeg","webp"], key="settings_logo")
        if logo_upload is not None:
            st.image(logo_upload, width=180)
            if st.button("💾 Lưu logo", type="primary", key="save_logo"):
                save_setting("logo", logo_upload.getvalue()); save_setting("logo_mime", logo_upload.type); st.success("Đã lưu logo."); st.rerun()
        if logo_bytes:
            st.image(logo_bytes, width=180, caption="Logo hiện tại")
            if st.button("🗑️ Xoá logo", key="delete_logo"):
                delete_setting("logo"); delete_setting("logo_mime"); st.rerun()
        st.divider()
        st.subheader("Hình nền website")
        bg_upload = st.file_uploader("Tải hình nền", type=["png","jpg","jpeg","webp"], key="settings_background")
        if bg_upload is not None:
            st.image(bg_upload, use_container_width=True)
            if st.button("💾 Lưu hình nền", type="primary", key="save_background"):
                save_setting("background", bg_upload.getvalue()); save_setting("background_mime", bg_upload.type); st.success("Đã lưu hình nền."); st.rerun()
        if bg_bytes:
            st.image(bg_bytes, use_container_width=True, caption="Hình nền hiện tại")
            if st.button("🗑️ Khôi phục nền hồng mặc định", key="delete_background"):
                delete_setting("background"); delete_setting("background_mime"); st.rerun()
        st.divider()
        st.subheader("Màu giao diện")
        st.color_picker("Màu chủ đạo", value="#D85B78")
        st.info("🎀 Bản này dùng hồng pastel làm màu chủ đạo.")
    with tab2:
        st.subheader("Preview")
        c1,c2 = st.columns([1,2])
        with c1:
            if logo_bytes: st.image(logo_bytes, width=180)
            else: st.markdown("### ✦ KOC")
        with c2:
            st.markdown("<div style='background:rgba(255,255,255,.9);padding:24px;border-radius:20px;border:1px solid #f0d6dd;'><h3 style='color:#8d3048;margin-top:0'>Tổng quan số liệu TAP/Booking</h3><p>Giao diện hồng pastel • sidebar menu dạng thu gọn • icon đồng bộ.</p></div>", unsafe_allow_html=True)
